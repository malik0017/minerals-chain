from tests.conftest import Client


def test_security_checklist_fix_and_env(admin, db):
    from app.core import system_settings as ss
    ss.set_setting(db, "allow_impersonation", True, None)
    db.commit()
    ss._CACHE.clear()
    r = admin.post("/admin/control-center/security/fix", {"item": "impersonation"})
    assert "msg=" in r.headers["location"]
    db.expire_all()
    assert ss.get_setting(db, "allow_impersonation") is False
    assert "error=" in admin.post("/admin/control-center/security/fix", {"item": "nope"}).headers["location"]
    env = admin.get("/admin/control-center/security/env")
    assert env.status_code == 200 and "restart" in env.text


def test_rate_limit_memory_fallback():
    import pytest
    from app.core import rate_limit
    from app.core.exceptions import RateLimitExceededException
    rate_limit.reset("t", "x")
    rate_limit.check("t", "x", max_attempts=2, window_seconds=60)
    rate_limit.check("t", "x", max_attempts=2, window_seconds=60)
    with pytest.raises(RateLimitExceededException):
        rate_limit.check("t", "x", max_attempts=2, window_seconds=60)
    assert rate_limit.backend()["name"] in ("memory", "redis")


def test_file_encryption_roundtrip(monkeypatch, tmp_path):
    from cryptography.fernet import Fernet
    from app.core.config import settings
    from app.services import file_crypto
    monkeypatch.setattr(settings, "FILE_ENCRYPTION_KEY", Fernet.generate_key().decode())
    p = tmp_path / "x.pdf"
    file_crypto.write(p, b"%PDF-1.4 secret")
    raw = p.read_bytes()
    assert raw.startswith(file_crypto.MAGIC) and b"secret" not in raw
    assert file_crypto.read(p) == b"%PDF-1.4 secret"
    monkeypatch.setattr(settings, "FILE_ENCRYPTION_KEY", Fernet.generate_key().decode())
    import pytest
    with pytest.raises(file_crypto.FileCryptoError):
        file_crypto.read(p)


def test_backup_run_and_verify(admin, db, monkeypatch, tmp_path):
    from app.core.config import settings
    from app.models.backup import BackupRun
    from app.services import backup_service
    monkeypatch.setattr(settings, "BACKUP_DIR", str(tmp_path))
    r = admin.post("/admin/backups/run")
    assert "msg=" in r.headers["location"], r.headers["location"]
    db.expire_all()
    run = db.query(BackupRun).order_by(BackupRun.started_at.desc()).first()
    assert run.status == "success" and backup_service.verify(run.name)["ok"]
    assert admin.get(f"/admin/backups/{run.name}/manifest.json").status_code == 200
    assert admin.get("/admin/backups/../etc/db.dump").status_code in (404, 307, 308)
    assert admin.get("/admin/backups").status_code == 200


def test_inventory_receipt_transfer_adjust_and_auto_issue(db):
    from decimal import Decimal
    from app.models.inventory import InventoryMovement
    from app.models.user import User
    from app.services import inventory_service as inv, order_service
    from tests.test_brd_completion import make_order
    from app.core import system_settings as ss
    ss.set_setting(db, "inventory_auto_issue", True, None)
    ss.set_setting(db, "allow_negative_stock", True, None)
    db.commit()
    seller = db.query(User).filter_by(email="seller@mineralstest.com").one()
    c = seller.company
    wh1 = inv.create_warehouse(db, c, seller, {"name_en": "Test Yard A"})
    wh2 = inv.create_warehouse(db, c, seller, {"name_en": "Test Yard B"})
    order = make_order(db)
    from app.models.product import Product
    p = db.query(Product).filter_by(seller_company_id=c.id, status="verified").first()
    order.product_id = p.id
    db.commit()
    inv.record_receipt(db, c, seller, {"product_id": str(p.id), "warehouse_id": str(wh1.id), "quantity": "500", "unit_cost_sar": "80"})
    inv.record_transfer(db, c, seller, {"product_id": str(p.id), "warehouse_id": str(wh1.id), "to_warehouse_id": str(wh2.id), "quantity": "100"})
    import pytest
    with pytest.raises(inv.InventoryError):
        inv.record_adjustment(db, c, seller, {"product_id": str(p.id), "warehouse_id": str(wh1.id), "quantity": "-3", "note": ""})
    inv.record_adjustment(db, c, seller, {"product_id": str(p.id), "warehouse_id": str(wh1.id), "quantity": "-3", "note": "Count variance"})
    before = inv.on_hand(db, c.id, p.id)
    order_service.mark_shipped(db, order)
    issue = db.query(InventoryMovement).filter_by(order_id=order.id, movement_type="sale_issue").one()
    assert inv.on_hand(db, c.id, p.id) == before + issue.quantity and issue.quantity < 0
    page = Client("seller@mineralstest.com").get("/seller/inventory")
    assert page.status_code == 200 and issue.reference in page.text


def test_erp_export_odoo_and_sap(db):
    import io
    import json
    import zipfile
    from app.models.company import Company, CompanyRole
    from app.models.user import User
    from app.services import erp_export_service as erp, private_files
    from tests.test_brd_completion import make_order

    make_order(db)
    admin = db.query(User).filter(User.role == "admin").first()
    odoo = erp.build(db, target="odoo", datasets=list(erp.DATASETS), start=None, end=None, scope=None, user=admin)
    data, intact = private_files.read_verified(odoo.file_path, odoo.file_sha256)
    z = zipfile.ZipFile(io.BytesIO(data))
    assert intact and "manifest.json" in z.namelist() and "res.partner.csv" in z.namelist()
    assert json.loads(z.read("manifest.json"))["counts"]["orders"] >= 1

    seller = db.query(Company).filter(Company.role == CompanyRole.SELLER).first()
    sap = erp.build(db, target="sap_b1", datasets=list(erp.DATASETS), start=None, end=None, scope=seller, user=admin)
    names = zipfile.ZipFile(io.BytesIO(private_files.read_verified(sap.file_path, sap.file_sha256)[0])).namelist()
    assert "invoices" not in sap.datasets and "BusinessPartners.json" in names and "Items.json" in names

    import pytest
    with pytest.raises(erp.ErpExportError):
        erp.build(db, target="xero", datasets=["orders"], start=None, end=None, scope=None, user=admin)


def test_shipment_dispatch_tracking_and_delivery(db):
    import pytest
    from app.models.order import OrderStatus
    from app.models.shipment import ShipmentEvent
    from app.models.user import User
    from app.services import shipment_service as svc
    from tests.test_brd_completion import make_order

    order = make_order(db)
    seller = db.query(User).filter(User.company_id == order.seller_company_id).first()
    with pytest.raises(svc.ShipmentError):
        svc.dispatch(db, order, {"gross_weight_t": "10", "tare_weight_t": "20"}, seller)
    db.rollback()
    s = svc.dispatch(db, order, {"transport_mode": "truck", "carrier_name": "Hala Transport", "vehicle_plate": "1234 RJA",
                                 "gross_weight_t": "56.5", "tare_weight_t": "16.5"}, seller)
    db.refresh(order)
    assert order.status == OrderStatus.IN_TRANSIT and s.reference.startswith("SHP") and s.net_weight_t == 40
    svc.add_update(db, s, {"status": "at_checkpoint", "location": "Al Khurmah"}, seller)
    svc.add_update(db, s, {"status": "delivered", "received_net_t": "39.8"}, seller)
    db.refresh(order)
    db.refresh(s)
    assert order.status == OrderStatus.DELIVERED and s.status == "delivered" and float(s.variance_t) == pytest.approx(-0.2)
    assert db.query(ShipmentEvent).filter(ShipmentEvent.shipment_id == s.id).count() == 3
    with pytest.raises(svc.ShipmentError):
        svc.add_update(db, s, {"status": "in_transit"}, seller)
    st = svc.stats(db, seller_id=order.seller_company_id)
    assert st["delivered"] >= 1


def test_company_document_versions_review_and_reminders(db):
    from datetime import date, timedelta
    import pytest
    from app.models.company import Company, CompanyRole
    from app.models.user import User
    from app.services import company_document_service as docs

    pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
    seller = db.query(Company).filter(Company.role == CompanyRole.SELLER, Company.status == "approved").first()
    user = seller.users[0]
    admin = db.query(User).filter(User.role == "admin").first()
    soon = (date.today() + timedelta(days=5)).isoformat()
    with pytest.raises(docs.DocumentError):
        docs.upload(db, seller, user, {"doc_type": "vat_cert", "issued_on": soon, "expires_on": "2020-01-01"}, "vat.pdf", pdf)
    v1 = docs.upload(db, seller, user, {"doc_type": "zakat_cert", "expires_on": soon}, "zakat.pdf", pdf)
    v2 = docs.upload(db, seller, user, {"doc_type": "zakat_cert", "expires_on": soon}, "zakat2.pdf", pdf)
    db.refresh(v1)
    assert v2.version == v1.version + 1 and v2.is_current and not v1.is_current
    with pytest.raises(docs.DocumentError):
        docs.review(db, v2, admin, False, "no")
    docs.review(db, v2, admin, True)
    assert docs.health(v2) == "expiring"
    row = next(r for r in docs.checklist(db, seller) if r["type"] == "zakat_cert")
    assert row["doc"].id == v2.id
    assert docs.send_expiry_reminders(db) >= 1
    db.refresh(v2)
    assert v2.last_reminder_days == 7
    assert docs.read(v2)[1]


def test_insights_tabs_and_xlsx(db, admin):
    import io
    from datetime import date, timedelta
    from openpyxl import load_workbook
    from app.services import insights_service as ins

    d0, d1 = date.today() - timedelta(days=365), date.today()
    for tab in ins.TABS:
        data = ins.build(db, tab, d0, d1)
        assert len(data["kpis"]) == 4 and all(len(r) == len(data["columns"]) for r in data["rows"])
        assert admin.get(f"/admin/reports/insights?tab={tab}").status_code == 200
    wb = load_workbook(io.BytesIO(ins.workbook(db, d0, d1)))
    assert len(wb.sheetnames) == len(ins.TABS) + 1
    r = admin.get("/admin/reports/insights.xlsx?tab=funnel")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_monitoring_metrics_errors_and_jobs(db, admin):
    from app.models.monitoring import ErrorEvent, JobRun, RequestStat
    from app.services import job_service, monitoring_service as mon

    admin.get("/admin/control-center")
    admin.get("/admin/monitoring")
    mon.flush(db)
    assert db.query(RequestStat).filter(RequestStat.route == "/admin/control-center").count() >= 1
    import uuid
    marker = f"boom-{uuid.uuid4().hex}"
    try:
        raise RuntimeError(marker)
    except RuntimeError as exc:
        mon.record_error({"method": "GET", "path": "/x", "headers": []}, exc)
    assert db.query(ErrorEvent).filter(ErrorEvent.message == marker).count() == 1
    db.query(ErrorEvent).filter(ErrorEvent.message == marker).delete()
    db.commit()
    run = job_service.run(db, "quotation_expiry", trigger="manual")
    assert run.status == "success" and run.finished_at
    assert "quotation_expiry" not in job_service.due(db)
    ov = mon.overview(db, 24)
    assert ov["requests"] >= 2 and len(ov["labels"]) == 24
    r = admin.post("/admin/monitoring/jobs/metrics_maintenance/run")
    assert r.status_code == 303 and db.query(JobRun).filter(JobRun.job == "metrics_maintenance").count() >= 1
    assert admin.get("/admin/monitoring").status_code == 200



def test_credential_checks_sandbox(db, admin):
    from app.models.company import Company, CompanyRole
    from app.services import credential_check_service as ccs

    assert ccs.name_similarity("Al-Faisal Minerals Trading Co.", "AL FAISAL MINERALS TRADING") > 0.9
    seller = db.query(Company).filter(Company.role == CompanyRole.SELLER, Company.status == "approved").first()
    buyer = db.query(Company).filter(Company.role == CompanyRole.BUYER).first()
    c = ccs.run_check(db, seller, "cr", None)
    assert c.result == ccs.sandbox_case(seller.cr_number) and c.mode == "sandbox"
    assert ccs.sandbox_case("1010000999") == "expired" and ccs.sandbox_case("1234567") == "verified"
    import pytest
    with pytest.raises(ccs.CredentialCheckError):
        ccs.run_check(db, buyer, "mining_license", None)
    r = admin.post(f"/admin/credential-checks/{seller.id}/cr", {"back": f"/admin/companies/{seller.id}"})
    assert r.status_code == 303 and f"/admin/companies/{seller.id}" in r.headers["location"]
    assert admin.get(f"/admin/companies/{seller.id}").status_code == 200
    assert admin.get("/admin/credential-checks").status_code == 200


def test_api_tokens_and_v1_endpoints(db):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models.user import User, UserRole
    from app.services import api_token_service as svc

    from app.models.company import ApprovalStatus, Company
    seller = (db.query(User).join(Company, Company.id == User.company_id)
              .filter(User.role == UserRole.SELLER, User.is_active.is_(True), Company.status == ApprovalStatus.APPROVED).first())
    pending = (db.query(User).join(Company, Company.id == User.company_id)
               .filter(Company.status == ApprovalStatus.PENDING, User.is_active.is_(True)).first())
    token, raw = svc.issue(db, seller, "pytest", ["read"], "30")
    if pending:
        ptoken, praw = svc.issue(db, pending, "pytest-pending", ["read"], "30")
        assert TestClient(app).get("/api/v1/me", headers={"Authorization": f"Bearer {praw}"}).status_code == 401
        svc.revoke(db, ptoken, pending)
    c = TestClient(app)
    assert c.get("/api/v1/me").status_code == 401
    h = {"Authorization": f"Bearer {raw}"}
    me = c.get("/api/v1/me", headers=h).json()
    assert me["role"] == "seller" and me["token"]["scopes"] == ["read"]
    for path in ("/api/v1/dashboard", "/api/v1/orders", "/api/v1/shipments", "/api/v1/rfqs", "/api/v1/listings",
                 "/api/v1/inventory", "/api/v1/documents", "/api/v1/notifications?unread=true"):
        assert c.get(path, headers=h).status_code == 200, path
    orders = c.get("/api/v1/orders?size=5", headers=h).json()["items"]
    for o in orders:
        assert o["role"] == "seller" and (o["counterparty"] is None) == (not o["identity_revealed"])
    ships = c.get("/api/v1/shipments", headers=h).json()["items"]
    if ships:
        assert c.post(f"/api/v1/shipments/{ships[0]['id']}/events", json={"status": "in_transit"}, headers=h).status_code == 403
    svc.revoke(db, token, seller)
    assert c.get("/api/v1/me", headers=h).status_code == 401
    assert c.get("/manifest.webmanifest").json()["short_name"] == "Minerals Chain"
    sw = c.get("/sw.js")
    assert sw.status_code == 200 and "serviceWorker" not in sw.text and "caches" in sw.text
    assert c.get("/offline").status_code == 200


def test_api_token_ui_flow(admin):
    r = admin.post("/account/api-tokens", {"name": "UI token", "scopes": "read", "expiry": "30"})
    assert r.status_code == 303 and "mc_new_token" in r.headers.get("set-cookie", "")
    page = admin.get("/account/api-tokens")
    assert page.status_code == 200 and "mc_" in page.text
    assert admin.get("/admin/api-tokens").status_code == 200

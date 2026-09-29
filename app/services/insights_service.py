"""
app/services/insights_service.py
"""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import and_, func
from sqlalchemy.orm import Session

from app.models.company import Company, CompanyRole
from app.models.dispute import Dispute
from app.models.order import Order, OrderStatus
from app.models.quotation import Quotation, QuotationStatus
from app.models.rfq import RFQ
from app.models.shipment import Shipment
from app.models.subscription import Subscription, SubscriptionCharge, SubscriptionStatus
from app.models.verification import VerificationRequest, VerificationStatus
from app.services.analytics_service import _months_between

TABS = {
    "prices": ("bi-graph-up", "Price trends"),
    "funnel": ("bi-funnel", "RFQ funnel"),
    "sellers": ("bi-shop", "Seller performance"),
    "buyers": ("bi-bag-check", "Buyer activity"),
    "labs": ("bi-eyedropper", "Lab turnaround"),
    "subscriptions": ("bi-stars", "Subscriptions"),
}
LIVE = (OrderStatus.CONFIRMED, OrderStatus.IN_TRANSIT, OrderStatus.DELIVERED, OrderStatus.INVOICED,
        OrderStatus.COMPLETED, OrderStatus.RESOLVED, OrderStatus.DISPUTED)
DONE = (OrderStatus.COMPLETED, OrderStatus.RESOLVED)


def _f(v) -> float:
    return float(v or 0)


def _bounds(start: date, end: date):
    return (datetime.combine(start, datetime.min.time(), tzinfo=timezone.utc),
            datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=timezone.utc))


def _in(col, t0, t1):
    return and_(col >= t0, col < t1)


def _month(col):
    return func.to_char(col, "YYYY-MM")


def _pct(a, b):
    return round(100 * a / b, 1) if b else 0.0


def prices(db: Session, start: date, end: date) -> dict:
    t0, t1 = _bounds(start, end)
    keys = _months_between(start, end)
    rows = (db.query(RFQ.mineral_type, _month(Order.created_at), func.avg(Order.price_per_unit), func.count(Order.id),
                     func.sum(Order.quantity), func.min(Order.price_per_unit), func.max(Order.price_per_unit))
            .join(RFQ, RFQ.id == Order.rfq_id)
            .filter(_in(Order.created_at, t0, t1), Order.status.in_(LIVE), Order.price_per_unit.isnot(None))
            .group_by(RFQ.mineral_type, _month(Order.created_at)).all())
    agg = defaultdict(lambda: {"orders": 0, "qty": 0.0, "min": None, "max": None, "sum": 0.0, "months": {}})
    for mineral, m, avg, n, qty, lo, hi in rows:
        a = agg[mineral]
        a["orders"] += n
        a["qty"] += _f(qty)
        a["sum"] += _f(avg) * n
        a["min"] = _f(lo) if a["min"] is None else min(a["min"], _f(lo))
        a["max"] = _f(hi) if a["max"] is None else max(a["max"], _f(hi))
        a["months"][m] = round(_f(avg), 2)
    top = sorted(agg.items(), key=lambda kv: -kv[1]["orders"])
    series = []
    for mineral, a in top[:5]:
        last = None
        data = []
        for k in keys:
            last = a["months"].get(k, last)
            data.append(a["months"].get(k))
        series.append({"name": mineral, "data": data})
    table = []
    for mineral, a in top:
        ms = [a["months"][k] for k in keys if k in a["months"]]
        change = _pct(ms[-1] - ms[0], ms[0]) if len(ms) > 1 else 0.0
        table.append([mineral, a["orders"], round(a["qty"], 1), a["min"], round(a["sum"] / a["orders"], 2), a["max"],
                      ms[-1] if ms else None, change])
    return {"kpis": [("bi-gem", "Minerals traded", len(agg), "primary"),
                     ("bi-receipt", "Orders", sum(a["orders"] for a in agg.values()), "info"),
                     ("bi-box", "Volume (MT)", f"{sum(a['qty'] for a in agg.values()):,.0f}", "success"),
                     ("bi-arrow-up-right", "Biggest mover", max(table, key=lambda r: abs(r[7]))[0] if table else "—", "warning")],
            "months": keys, "series": series,
            "columns": ["Mineral", "Orders", "Volume (MT)", "Min price", "Avg price", "Max price", "Last month avg", "Change %"],
            "rows": table}


def funnel(db: Session, start: date, end: date) -> dict:
    t0, t1 = _bounds(start, end)
    rfqs = db.query(RFQ).filter(_in(RFQ.created_at, t0, t1)).all()
    ids = [r.id for r in rfqs]
    quotes = db.query(Quotation).filter(Quotation.rfq_id.in_(ids)).all() if ids else []
    orders = db.query(Order).filter(Order.rfq_id.in_(ids)).all() if ids else []
    by_rfq = defaultdict(list)
    for q in quotes:
        by_rfq[q.rfq_id].append(q)
    quoted = [r for r in rfqs if by_rfq[r.id]]
    ordered = {o.rfq_id for o in orders}
    confirmed = [o for o in orders if o.confirmed_at]
    done = [o for o in orders if o.status in DONE]
    hours = [(min(q.created_at for q in by_rfq[r.id]) - r.created_at).total_seconds() / 3600 for r in quoted]
    per_mineral = defaultdict(lambda: [0, 0, 0, 0])
    for r in rfqs:
        m = per_mineral[r.mineral_type]
        m[0] += 1
        m[1] += bool(by_rfq[r.id])
        m[2] += len(by_rfq[r.id])
        m[3] += r.id in ordered
    mineral_rows = sorted([[k, v[0], _pct(v[1], v[0]), round(v[2] / v[0], 1), _pct(v[3], v[0])] for k, v in per_mineral.items()],
                          key=lambda r: -r[1])
    stages = [("RFQs posted", len(rfqs)), ("Received quotes", len(quoted)), ("Quote accepted", len(ordered)),
              ("Seller confirmed", len(confirmed)), ("Completed", len(done))]
    return {"kpis": [("bi-file-earmark-text", "RFQs", len(rfqs), "primary"),
                     ("bi-tags", "Avg quotes / RFQ", round(len(quotes) / len(rfqs), 1) if rfqs else 0, "info"),
                     ("bi-stopwatch", "Hours to first quote", round(sum(hours) / len(hours), 1) if hours else "—", "warning"),
                     ("bi-trophy", "RFQ → order", f"{_pct(len(ordered), len(rfqs))}%", "success")],
            "stages": [{"name": n, "value": v} for n, v in stages],
            "columns": ["Mineral", "RFQs", "Quoted %", "Avg quotes", "Converted %"],
            "by_mineral": [{"name": r[0], "value": r[1]} for r in mineral_rows[:10]], "rows": mineral_rows}


def sellers(db: Session, start: date, end: date) -> dict:
    t0, t1 = _bounds(start, end)
    q = dict(db.query(Quotation.seller_company_id, func.count()).filter(_in(Quotation.created_at, t0, t1))
             .group_by(Quotation.seller_company_id).all())
    won = dict(db.query(Quotation.seller_company_id, func.count())
               .filter(_in(Quotation.created_at, t0, t1), Quotation.status == QuotationStatus.ACCEPTED)
               .group_by(Quotation.seller_company_id).all())
    lead = dict(db.query(Quotation.seller_company_id, func.avg(Quotation.lead_time_days)).filter(_in(Quotation.created_at, t0, t1))
                .group_by(Quotation.seller_company_id).all())
    orders = (db.query(Order.seller_company_id, func.count(), func.coalesce(func.sum(Order.total_value_sar), 0))
              .filter(_in(Order.created_at, t0, t1), Order.status.in_(LIVE)).group_by(Order.seller_company_id).all())
    om = {sid: (n, _f(v)) for sid, n, v in orders}
    ships = db.query(Shipment).filter(_in(Shipment.dispatched_at, t0, t1), Shipment.status == "delivered").all()
    ontime = defaultdict(lambda: [0, 0])
    for s in ships:
        ontime[s.seller_company_id][1] += 1
        ontime[s.seller_company_id][0] += (not s.eta or s.delivered_at <= s.eta)
    disputes = dict(db.query(Order.seller_company_id, func.count(Dispute.id)).join(Order, Order.id == Dispute.order_id)
                    .filter(_in(Dispute.created_at, t0, t1)).group_by(Order.seller_company_id).all())
    names = dict(db.query(Company.id, Company.company_name).filter(Company.role == CompanyRole.SELLER).all())
    rows = []
    for sid in set(q) | set(om):
        n, gmv = om.get(sid, (0, 0.0))
        ot = ontime.get(sid)
        rows.append([names.get(sid, "—"), q.get(sid, 0), _pct(won.get(sid, 0), q.get(sid, 0)), n, round(gmv, 2),
                     round(_f(lead.get(sid)), 1), _pct(ot[0], ot[1]) if ot else None, disputes.get(sid, 0)])
    rows.sort(key=lambda r: -r[4])
    total_q = sum(q.values())
    return {"kpis": [("bi-shop", "Active sellers", len(rows), "primary"),
                     ("bi-tags", "Quotations", total_q, "info"),
                     ("bi-trophy", "Win rate", f"{_pct(sum(won.values()), total_q)}%", "success"),
                     ("bi-truck", "On-time delivery", f"{_pct(sum(v[0] for v in ontime.values()), sum(v[1] for v in ontime.values()))}%", "warning")],
            "top": [{"name": r[0], "value": r[4]} for r in rows[:10]],
            "columns": ["Seller", "Quotes", "Win %", "Orders", "GMV (SAR)", "Avg lead (days)", "On-time %", "Disputes"],
            "rows": rows}


def buyers(db: Session, start: date, end: date) -> dict:
    t0, t1 = _bounds(start, end)
    rfqs = dict(db.query(RFQ.buyer_company_id, func.count()).filter(_in(RFQ.created_at, t0, t1)).group_by(RFQ.buyer_company_id).all())
    orders = (db.query(Order.buyer_company_id, func.count(), func.coalesce(func.sum(Order.total_value_sar), 0), func.max(Order.created_at))
              .filter(_in(Order.created_at, t0, t1), Order.status.in_(LIVE)).group_by(Order.buyer_company_id).all())
    om = {b: (n, _f(v), last) for b, n, v, last in orders}
    disputes = dict(db.query(Order.buyer_company_id, func.count(Dispute.id)).join(Order, Order.id == Dispute.order_id)
                    .filter(_in(Dispute.created_at, t0, t1)).group_by(Order.buyer_company_id).all())
    names = dict(db.query(Company.id, Company.company_name).filter(Company.role == CompanyRole.BUYER).all())
    rows = []
    for bid in set(rfqs) | set(om):
        n, gmv, last = om.get(bid, (0, 0.0, None))
        rows.append([names.get(bid, "—"), rfqs.get(bid, 0), n, _pct(n, rfqs.get(bid, 0)), round(gmv, 2),
                     round(gmv / n, 2) if n else 0, disputes.get(bid, 0), last.strftime("%Y-%m-%d") if last else "—"])
    rows.sort(key=lambda r: -r[4])
    gmv = sum(r[4] for r in rows)
    top5 = sum(r[4] for r in rows[:5])
    return {"kpis": [("bi-bag-check", "Active buyers", len(rows), "primary"),
                     ("bi-cash-stack", "GMV (SAR)", f"{gmv:,.0f}", "success"),
                     ("bi-basket", "Avg order (SAR)", f"{gmv / max(1, sum(r[2] for r in rows)):,.0f}", "info"),
                     ("bi-pie-chart", "Top-5 share", f"{_pct(top5, gmv)}%", "warning")],
            "top": [{"name": r[0], "value": r[4]} for r in rows[:10]],
            "columns": ["Buyer", "RFQs", "Orders", "RFQ → order %", "GMV (SAR)", "Avg order (SAR)", "Disputes", "Last order"],
            "rows": rows}


def labs(db: Session, start: date, end: date) -> dict:
    t0, t1 = _bounds(start, end)
    reqs = db.query(VerificationRequest).filter(_in(VerificationRequest.created_at, t0, t1)).all()
    names = dict(db.query(Company.id, Company.company_name).filter(Company.role == CompanyRole.LAB).all())
    per = defaultdict(lambda: {"n": 0, "done": 0, "failed": 0, "open": 0, "days": [], "fees": 0.0, "priority": 0})
    buckets = {"≤ 3 days": 0, "4–7 days": 0, "8–14 days": 0, "> 14 days": 0}
    for r in reqs:
        p = per[r.lab_company_id]
        p["n"] += 1
        p["priority"] += bool(r.is_priority)
        p["fees"] += _f(r.fee_sar)
        if r.status in (VerificationStatus.COMPLETED, VerificationStatus.FAILED) and r.completed_at:
            d = (r.completed_at - r.created_at).total_seconds() / 86400
            p["days"].append(d)
            p["done" if r.status == VerificationStatus.COMPLETED else "failed"] += 1
            buckets["≤ 3 days" if d <= 3 else "4–7 days" if d <= 7 else "8–14 days" if d <= 14 else "> 14 days"] += 1
        else:
            p["open"] += 1
    rows = []
    for lid, p in per.items():
        closed = p["done"] + p["failed"]
        rows.append([names.get(lid, "—"), p["n"], p["open"], closed, _pct(p["done"], closed),
                     round(sum(p["days"]) / len(p["days"]), 1) if p["days"] else None, p["priority"], round(p["fees"], 2)])
    rows.sort(key=lambda r: -r[1])
    all_days = [d for p in per.values() for d in p["days"]]
    closed = sum(p["done"] + p["failed"] for p in per.values())
    return {"kpis": [("bi-eyedropper", "Requests", len(reqs), "primary"),
                     ("bi-stopwatch", "Avg turnaround (days)", round(sum(all_days) / len(all_days), 1) if all_days else "—", "warning"),
                     ("bi-patch-check", "Pass rate", f"{_pct(sum(p['done'] for p in per.values()), closed)}%", "success"),
                     ("bi-hourglass", "Open", sum(p["open"] for p in per.values()), "info")],
            "buckets": [{"name": k, "value": v} for k, v in buckets.items()],
            "turnaround": [{"name": r[0], "value": r[5] or 0} for r in rows],
            "columns": ["Lab", "Requests", "Open", "Closed", "Pass %", "Avg days", "Priority", "Fees (SAR)"],
            "rows": rows}


def subscriptions(db: Session, start: date, end: date) -> dict:
    keys = _months_between(start, end)
    charges = (db.query(func.to_char(SubscriptionCharge.period_start, "YYYY-MM"), SubscriptionCharge.tier, func.sum(SubscriptionCharge.amount_sar))
               .filter(SubscriptionCharge.period_start >= start, SubscriptionCharge.period_start <= end,
                       SubscriptionCharge.status.in_(["paid", "pending"]))
               .group_by(func.to_char(SubscriptionCharge.period_start, "YYYY-MM"), SubscriptionCharge.tier).all())
    tiers = sorted({t for _, t, _ in charges})
    grid = {t: {k: 0.0 for k in keys} for t in tiers}
    for m, t, v in charges:
        if m in grid[t]:
            grid[t][m] += _f(v)
    series = [{"name": t.title(), "data": [round(grid[t][k], 2) for k in keys]} for t in tiers]
    totals = [round(sum(grid[t][k] for t in tiers), 2) for k in keys]
    active = (db.query(Subscription.tier, func.count()).filter(Subscription.status.in_([SubscriptionStatus.ACTIVE, SubscriptionStatus.GRACE]))
              .group_by(Subscription.tier).all())
    t0, t1 = _bounds(start, end)
    churned = db.query(Subscription).filter(Subscription.status == SubscriptionStatus.EXPIRED, _in(Subscription.updated_at, t0, t1)).count()
    new = db.query(Subscription).filter(_in(Subscription.created_at, t0, t1)).count()
    status = dict(db.query(SubscriptionCharge.status, func.count()).filter(SubscriptionCharge.period_start >= start,
                                                                            SubscriptionCharge.period_start <= end)
                  .group_by(SubscriptionCharge.status).all())
    rows = [[k] + [round(grid[t][k], 2) for t in tiers] + [totals[i]] for i, k in enumerate(keys)]
    return {"kpis": [("bi-graph-up-arrow", "Billed last month (SAR)", f"{totals[-1]:,.0f}" if totals else 0, "success"),
                     ("bi-people", "Active subscriptions", sum(n for _, n in active), "primary"),
                     ("bi-person-plus", "New in period", new, "info"),
                     ("bi-person-dash", "Expired in period", churned, "danger")],
            "months": keys, "series": series,
            "active": [{"name": (t.value if hasattr(t, "value") else str(t)).title(), "value": n} for t, n in active],
            "charge_status": [{"name": k.title(), "value": v} for k, v in status.items()],
            "columns": ["Month"] + [t.title() for t in tiers] + ["Total"], "rows": rows}


BUILDERS = {"prices": prices, "funnel": funnel, "sellers": sellers, "buyers": buyers, "labs": labs, "subscriptions": subscriptions}


def build(db: Session, tab: str, start: date, end: date) -> dict:
    return BUILDERS.get(tab, prices)(db, start, end)


def workbook(db: Session, start: date, end: date, tabs: list[str] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Minerals Chain — analytics", f"{start} → {end}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    head_fill = PatternFill("solid", fgColor="0F766E")
    for key in tabs or list(TABS):
        data = build(db, key, start, end)
        summary.append([TABS[key][1]])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
        for _, label, value, _c in data["kpis"]:
            summary.append(["", label, value])
        ws = wb.create_sheet(TABS[key][1][:31].replace("&", "and"))
        ws.append(data["columns"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for r in data["rows"]:
            ws.append(r)
        for i, col in enumerate(data["columns"], 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in data["rows"][:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
        ws.freeze_panes = "A2"
        if data["rows"]:
            ws.auto_filter.ref = ws.dimensions
    summary.column_dimensions["B"].width = 32
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
